"""SIF 아카이브 → Neo4j 적재 도구.

사용법
  python ingest.py check                      Neo4j 연결 확인
  python ingest.py init                       제약조건·인덱스 만들기 (처음 한 번)
  python ingest.py load 파일.xlsx [파일2 ...]  엑셀을 읽어 적재 (같은 사례는 갱신, 새 사례는 추가)
  python ingest.py load 파일.xlsx --dry-run    Neo4j에 쓰지 않고 preview/ 폴더에 결과만 저장
  python ingest.py remap                      config/ 규칙을 바꾼 뒤 이미 올라간 사례 전체에 다시 적용
  python ingest.py derive                     파생 관계(LEADS_TO, MITIGATES)만 다시 계산
  python ingest.py terms                      config/aliases.csv(용어 사전)만 다시 적용
  python ingest.py prune                      config/scope.yaml 범위 밖 사례(예: 제조업)를 DB에서 지우고 다시 계산
  python ingest.py template                   정리 양식 엑셀(template/SIF_정리양식.xlsx) 만들기
"""
from __future__ import annotations

import argparse
import csv
import os
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from sif.read import read_workbook  # noqa: E402
from sif.transform import Config, transform, raw_from_props  # noqa: E402


def load_env():
    f = ROOT / '.env'
    if f.exists():
        for line in f.read_text(encoding='utf-8').splitlines():
            if '=' in line and not line.strip().startswith('#'):
                k, v = line.split('=', 1)
                os.environ.setdefault(k.strip(), v.strip())


def write_review(cases, out: Path):
    out.mkdir(parents=True, exist_ok=True)
    issues = [i for c in cases for i in c.issues]
    with open(out / 'review_needed.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['사례ID', '항목', '값', '사유'])
        w.writerows(issues)
    agg = Counter((i[1], i[2]) for i in issues)
    with open(out / 'review_summary.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['항목', '값', '사례 수'])
        for (k, v), n in agg.most_common():
            w.writerow([k, v, n])
    return issues, agg


def write_preview(cases, out: Path):
    out.mkdir(parents=True, exist_ok=True)
    with open(out / 'cases_mapped.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['id', 'code', '업종구분', '원본 위치', '재해종류', '재해종류 출처', '상위분류', '사례별 보정',
                    '공식 발생형태', '기인물', '기인물(상위)', '기인물 분류', '유발요인', '요인 분류', '대책 수',
                    '통제위계', '연도', '월', '시기', '추락 높이(m)', '높이 구간', '사망', '재해개요'])
        for c in cases:
            w.writerow([c.id, c.props['code'], c.props['sector'], f"{c.props['sheet']} {c.props['row']}행", c.at,
                        c.at_source, c.at_group, c.classified_as or '', c.at_official, c.co, c.co_base,
                        c.co_category, ', '.join(h['name'] for h in c.hazards),
                        ', '.join(sorted({h['category'] for h in c.hazards})), len(c.measures),
                        ', '.join(sorted({lv for m in c.measures for lv in m['levels']})), c.year or '', c.month or '',
                        c.period or '', c.height or '', c.band or '', 'Y' if c.props['fatal'] else '', c.props['summary']])


def report_terms(missing, out: Path):
    if not missing:
        print('  용어 사전: 모든 용어가 개념에 연결되었습니다.')
        return
    with open(out / 'terms_unmatched.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, list(missing[0].keys()))
        w.writeheader()
        w.writerows(missing)
    print(f'  용어 사전: 연결 못 한 용어 {len(missing)}개 → {out}/terms_unmatched.csv')


def write_pattern_preview(cases, cfg, out: Path):
    from sif.patterns import build_patterns, coverage
    pats = build_patterns([c.props for c in cases], cfg.patterns)
    by = {c.id: c for c in cases}
    with open(out / 'patterns.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['패턴', '사례 수', '업종', '주요 작업', '대표 사례 1', '대표 사례 1 개요', '대표 사례 2', '대표 사례 2 개요'])
        for p in pats:
            reps = [by[r] for r in p['reps']]
            row = [p['name'], p['size'], '/'.join(p['sectors']), ', '.join(p['topWorks'])]
            for k in range(2):
                row += [reps[k].props['code'], reps[k].props['summary']] if k < len(reps) else ['', '']
            w.writerow(row)
    print(f'  사고 패턴: {coverage(pats, len(cases))} → {out}/patterns.csv')


def check_terms_local(cases, cfg):
    """dry-run용: Neo4j 없이 변환 결과의 이름들로 용어 사전이 맞는지 확인합니다."""
    names = {
        'CausalObject': {c.co for c in cases} | {c.co_base for c in cases},
        'ObjectCategory': {c.co_category for c in cases}, 'AccidentType': {c.at for c in cases},
        'AccidentTypeGroup': {c.at_group for c in cases}, 'HazardFactor': {h['name'] for c in cases for h in c.hazards},
        'HazardCategory': {h['category'] for c in cases for h in c.hazards},
        'ControlLevel': {lv for c in cases for m in c.measures for lv in m['levels']},
        'WorkType': {c.context.get('wt', '') for c in cases}, 'Work': {c.context.get('w', '') for c in cases},
        'UnitWork': {c.context.get('uw', '') for c in cases}, 'Period': {c.period for c in cases if c.period},
        'Year': {f'{c.year}년' for c in cases if c.year},
        'HeightBand': {c.band for c in cases if c.band} | {c.band_parent for c in cases if c.band_parent},
    }
    missing, hits = [], Counter()
    for a in cfg.aliases:
        pool = names.get(a['대상 라벨'], set())
        n = sum(1 for x in pool if x and (a['대상 이름'] in x if a['일치 방식'] == 'contains' else x == a['대상 이름']))
        hits[a['용어']] += n
        if n == 0:
            missing.append({**a, '사유': '대상 노드 없음'})
    return missing, hits


def report(cases, issues, agg):
    print(f'  사례 {len(cases):,}건  (제조업 등 {sum(c.props["sectorCode"] == "M" for c in cases):,}'
          f' / 건설업 {sum(c.props["sectorCode"] == "C" for c in cases):,})')
    print(f'  검토 필요 {len(issues):,}건')
    for (k, v), n in agg.most_common(8):
        print(f'    - {k}: {v[:40]} ({n}건)')


def cmd_load(args, cfg):
    rows, ids = [], set()
    for p in args.files:
        r, warns = read_workbook(p, cfg.columns, archive=args.archive)
        print(f'{p}')
        for w in warns:
            print('  ' + w)
        for x in r:
            if x['id'] in ids:
                continue
            ids.add(x['id'])
            x['_file'] = p
            rows.append(x)
    skipped = Counter(r['sector'] for r in rows if r['sector'] not in cfg.sectors)
    rows = [r for r in rows if r['sector'] in cfg.sectors]
    if skipped:
        names = {'C': '건설업', 'M': '제조업 등'}
        print('  적재 범위(config/scope.yaml) 밖이라 제외: ' + ', '.join(f'{names[k]} {v:,}건' for k, v in skipped.items()))
    cases = [transform(r, cfg) for r in rows]
    out = Path(args.out)
    issues, agg = write_review(cases, out)
    write_preview(cases, out)
    report(cases, issues, agg)
    print(f'  미리보기와 검토 목록: {out}/cases_mapped.csv, {out}/review_needed.csv')
    write_pattern_preview(cases, cfg, out)
    if args.dry_run:
        missing, hits = check_terms_local(cases, cfg)
        amb = [(t, n) for t, n in hits.items() if n > 3]
        print(f'  용어 사전: {len(cfg.aliases)}줄, 연결 못 한 줄 {len(missing)}개, 여러 개념을 가리키는 용어 {len(amb)}개'
              + (f' (예: {", ".join(f"{t} {n}개" for t, n in amb[:3])})' if amb else ''))
        report_terms(missing, out)
        print('dry-run: Neo4j에는 쓰지 않았습니다.')
        return
    from sif.neo import Writer, driver_from_env
    with driver_from_env() as d:
        w = Writer(d)
        w.init()
        existing = w.existing_ids([c.id for c in cases])
        new = len(cases) - len(existing)
        for p in args.files:
            src = Path(p).name
            part = [c for c, r in zip(cases, rows) if r['_file'] == p]
            w.write_cases(part, src)
            w.record_run(src, [c.id for c in part], sum(c.id not in existing for c in part))
        w.derive()
        removed = w.cleanup()
        pats = w.write_patterns(cfg.patterns)
        missing = w.write_terms(cfg.aliases)
        print(f'적재 완료: 새 사례 {new:,}건, 갱신 {len(existing):,}건, 정리된 빈 노드 {removed}개')
        from sif.patterns import coverage
        print('  사고 패턴: ' + coverage(pats, sum(r['n'] for r in w.summary() if r['label'] == 'AccidentCase')))
        report_terms(missing, Path(args.out))
        for r in w.summary():
            print(f"  {r['label']:<20}{r['n']:>8,}")


def cmd_remap(args, cfg):
    from sif.neo import Writer, driver_from_env
    with driver_from_env() as d:
        w = Writer(d)
        props = [p for p in w.all_case_props() if p.get('sectorCode') in cfg.sectors]
        extra = w.out_of_scope(cfg.sectors)
        if extra:
            print(f'  적재 범위 밖 사례 {extra:,}건이 DB에 남아 있습니다. 지우려면 python ingest.py prune')
        print(f'Neo4j에서 사례 {len(props):,}건을 읽었습니다. 현재 config/ 규칙으로 다시 매핑합니다.')
        cases = [transform(raw_from_props(p), cfg) for p in props]
        issues, agg = write_review(cases, Path(args.out))
        write_preview(cases, Path(args.out))
        report(cases, issues, agg)
        if args.dry_run:
            print('dry-run: Neo4j에는 쓰지 않았습니다.')
            return
        w.write_cases(cases, None)
        w.derive()
        print(f'다시 매핑 완료. 정리된 빈 노드 {w.cleanup()}개')
        from sif.patterns import coverage
        print('  사고 패턴: ' + coverage(w.write_patterns(cfg.patterns), len(cases)))
        report_terms(w.write_terms(cfg.aliases), Path(args.out))


def cmd_template(args, cfg):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = '사례'
    head = ['사례ID', '업종구분', '연번', '산재업종(대분류)', '산재업종(중분류)', '산재업종(소분류)', '고위험작업·상황',
            '공종', '작업명', '단위작업명', '재해종류', '재해개요', '기인물', '재해유발요인', '위험성 감소대책(예시)']
    ws.append(head)
    ws.append(['', '제조업 등', 1, '제조업', '화학및고무제품제조업', '화학비료제조업', '비정형 작업(점검)', '', '', '', '',
               '2019년 1월경 OO사업장 생산동 지붕에서 재해자가 현장 점검 중 아래로 떨어져 사망', '지붕·대들보(슬레이트)',
               '『밟고 있던 구조물(또는 자재) 등의 파손·탈락 등』으로 넘어지거나 떨어짐 위험',
               '▶ 지붕으로 이동할 수 있는 장소에 방책 설치\n▶ 30센티미터 이상의 발판 설치 또는 안전대 착용'])
    ws.append(['', '건설업', 2, '', '', '', '', '1. 토공사', '1.1 굴착 작업', '1.1.1 굴착 장비반입', '전도',
               '2020년 09월경 굴삭기를 화물자동차에서 내리던 중 넘어지는 굴삭기에 깔려 사망', '굴착기(백호우)',
               '굴착기를 화물차에서 내리는 중 전도되는 굴착기에 깔림',
               '▶ 충분한 길이·폭·강도의 발판 사용\n▶ 평탄한 지반에서 작업계획서에 따라 하역'])
    bold, fill = Font(name='Arial', bold=True), PatternFill('solid', fgColor='FFF2C230')
    for c in ws[1]:
        c.font, c.fill, c.alignment = bold, fill, Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name='Arial', color='FF0000FF')
            c.alignment = Alignment(wrap_text=True, vertical='top')
    widths = [14, 11, 6, 12, 18, 16, 18, 16, 18, 20, 10, 50, 22, 44, 50]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wdt
    ws.freeze_panes = 'A2'
    dv = DataValidation(type='list', formula1='"제조업 등,건설업"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add('B2:B5000')
    g = wb.create_sheet('작성 안내')
    guide = [
        ['항목', '설명'],
        ['작성 위치', '사례 시트 2행부터 한 행에 사례 하나. 파란 글씨 2행은 예시이니 지우고 쓰세요.'],
        ['업종구분', '제조업 등 / 건설업 중 선택 (필수)'],
        ['사례ID', '비워 두면 재해개요·기인물·유발요인으로 고유 ID를 자동으로 만듭니다.'],
        ['제조업 등', '산재업종 대·중·소분류와 고위험작업·상황을 채우고, 공종·작업명·단위작업명은 비웁니다.'],
        ['건설업', '공종·작업명·단위작업명과 재해종류를 채우고, 산재업종과 고위험작업·상황은 비웁니다.'],
        ['재해종류', '비워 두면 재해개요·유발요인 문장에서 자동 추출합니다 (config/rules_accident_type.yaml).'],
        ['위험성 감소대책', '대책마다 앞에 ▶ 를 붙이세요. ▶가 없으면 줄바꿈으로 나눕니다.'],
        ['공단 원본', '공단 아카이브 원본 엑셀은 이 양식으로 옮기지 않고 그대로 올려도 됩니다.'],
    ]
    for r in guide:
        g.append(r)
    for c in g[1]:
        c.font = bold
    g.column_dimensions['A'].width = 16
    g.column_dimensions['B'].width = 90
    out = ROOT / 'template' / 'SIF_정리양식.xlsx'
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f'양식을 만들었습니다: {out}')


def main():
    load_env()
    ap = argparse.ArgumentParser(description='SIF 아카이브 → Neo4j 적재 도구')
    sub = ap.add_subparsers(dest='cmd', required=True)
    for name in ('check', 'init', 'derive', 'template', 'terms', 'prune'):
        sub.add_parser(name)
    p = sub.add_parser('load')
    p.add_argument('files', nargs='+')
    p.add_argument('--dry-run', action='store_true')
    p.add_argument('--archive', help='아카이브 버전 이름 (기본: 파일 이름)')
    p.add_argument('--out', default=str(ROOT / 'preview'))
    p = sub.add_parser('remap')
    p.add_argument('--dry-run', action='store_true')
    p.add_argument('--out', default=str(ROOT / 'preview'))
    ap.add_argument('--config', default=str(ROOT / 'config'))
    args = ap.parse_args()
    cfg = Config.load(args.config)
    if args.cmd == 'load':
        cmd_load(args, cfg)
    elif args.cmd == 'remap':
        cmd_remap(args, cfg)
    elif args.cmd == 'template':
        cmd_template(args, cfg)
    else:
        from sif.neo import Writer, driver_from_env
        with driver_from_env() as d:
            w = Writer(d)
            if args.cmd == 'check':
                d.verify_connectivity()
                print('연결 성공')
                for r in w.summary():
                    print(f"  {r['label']:<20}{r['n']:>8,}")
                extra = w.out_of_scope(cfg.sectors)
                if extra:
                    print(f'  적재 범위(config/scope.yaml: {cfg.sectors}) 밖 사례 {extra:,}건 → python ingest.py prune 으로 정리')
            elif args.cmd == 'prune':
                n = w.prune(cfg.sectors)
                print(f'적재 범위 밖 사례 {n:,}건을 지웠습니다. 파생 관계·패턴·용어를 다시 계산합니다.')
                w.derive()
                print(f'  정리된 빈 노드 {w.cleanup()}개')
                from sif.patterns import coverage
                kept = sum(r['n'] for r in w.summary() if r['label'] == 'AccidentCase')
                print('  사고 패턴: ' + coverage(w.write_patterns(cfg.patterns), kept))
                report_terms(w.write_terms(cfg.aliases), ROOT / 'preview')
            elif args.cmd == 'init':
                w.init()
                print('제약조건·인덱스를 만들었습니다.')
            elif args.cmd == 'derive':
                w.derive()
                print('파생 관계를 다시 계산했습니다.')
            elif args.cmd == 'terms':
                out = ROOT / 'preview'
                out.mkdir(exist_ok=True)
                report_terms(w.write_terms(cfg.aliases), out)


if __name__ == '__main__':
    main()
